#!/usr/bin/env python3
"""
Converts a handful of markdown docs (test plan, risk matrix, release
readiness checklist, Task 2 strategy) to .docx, since reviewers/stakeholders
often want these as Word documents rather than markdown.

Supports exactly the markdown subset this repo's docs actually use:
headings (#/##/###), bold/italic/inline-code, bullet and numbered lists,
GitHub-style task list checkboxes (- [ ]), tables, and --- horizontal rules.
It is not a general-purpose markdown parser - keep source docs within that
subset.

Usage: python3 scripts/md_to_docx.py docs/TEST_PLAN.md docs/RISK_MATRIX.md ...
Each output file is written alongside its source, with a .docx extension.
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

INLINE_TOKEN = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`|\[[^\]]+?\]\([^)]+?\))")
HEADING_COLOR = RGBColor(0x0B, 0x1F, 0x3A)
ACCENT_COLOR = RGBColor(0x0B, 0x5F, 0xFF)


def is_table_separator(line: str) -> bool:
    """Matches a markdown header-separator row, e.g. `|---|:--:|---|`."""
    line = line.strip()
    if not line.startswith("|"):
        return False
    cells = [c.strip() for c in line.strip("|").split("|")]
    return bool(cells) and all(re.match(r"^:?-+:?$", c) for c in cells)


def add_inline_runs(paragraph, text):
    """Splits `text` on markdown inline tokens and appends styled runs."""
    for token in INLINE_TOKEN.split(text):
        if not token:
            continue
        if token.startswith("**") and token.endswith("**"):
            paragraph.add_run(token[2:-2]).bold = True
        elif token.startswith("*") and token.endswith("*"):
            paragraph.add_run(token[1:-1]).italic = True
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Courier New"
            run.font.size = Pt(10)
        elif token.startswith("[") and "](" in token:
            visible = token[1:token.index("](")]
            paragraph.add_run(visible)
        else:
            paragraph.add_run(token)


def set_cell_shading(cell, hex_color):
    shd = cell._element.get_or_add_tcPr()
    el = shd.makeelement(qn("w:shd"), {qn("w:fill"): hex_color})
    shd.append(el)


def convert(md_path: Path, out_path: Path):
    lines = md_path.read_text().splitlines()
    doc = Document()

    # Base font for the whole document.
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)

    i = 0
    title_done = False
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if stripped == "---":
            doc.add_paragraph().add_run("").add_break()
            i += 1
            continue

        # Table block: a header row, a separator row (|---|---|), then data rows.
        if stripped.startswith("|") and i + 1 < len(lines) and is_table_separator(lines[i + 1]):
            header_cells = [c.strip() for c in stripped.strip("|").split("|")]
            i += 2
            data_rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                data_rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            table = doc.add_table(rows=1, cols=len(header_cells))
            table.style = "Light Grid Accent 1"
            for c, text in zip(table.rows[0].cells, header_cells):
                p = c.paragraphs[0]
                add_inline_runs(p, text)
                for run in p.runs:
                    run.bold = True
                set_cell_shading(c, "1F3864")
                for run in p.runs:
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            for row in data_rows:
                cells = table.add_row().cells
                for c, text in zip(cells, row):
                    add_inline_runs(c.paragraphs[0], text)
            doc.add_paragraph()
            continue

        # Headings.
        m = re.match(r"^(#{1,3})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            text = m.group(2)
            if level == 1 and not title_done:
                h = doc.add_heading("", level=0)
                title_done = True
            else:
                h = doc.add_heading("", level=level)
            add_inline_runs(h, text)
            for run in h.runs:
                run.font.color.rgb = HEADING_COLOR
            i += 1
            continue

        # Checkbox list item: - [ ] text. Uses "List Paragraph" (indented, no
        # auto-bullet glyph) rather than "List Bullet", since the manual
        # checkbox character below is the marker - "List Bullet" would add
        # its own bullet dot in front of it too.
        m = re.match(r"^[-*]\s+\[( |x|X)\]\s+(.*)$", stripped)
        if m:
            checked = m.group(1).lower() == "x"
            text = m.group(2)
            p = doc.add_paragraph(style="List Paragraph")
            box = p.add_run("☑ " if checked else "☐ ")
            box.font.size = Pt(11)
            add_inline_runs(p, text)
            i += 1
            continue

        # Numbered list item: 1. text
        m = re.match(r"^\d+\.\s+(.*)$", stripped)
        if m:
            p = doc.add_paragraph(style="List Number")
            add_inline_runs(p, m.group(1))
            i += 1
            continue

        # Nested bullet (2-space indented under a bullet above).
        m = re.match(r"^\s{2,}[-*]\s+(.*)$", line)
        if m:
            p = doc.add_paragraph(style="List Bullet 2")
            add_inline_runs(p, m.group(1))
            i += 1
            continue

        # Plain bullet: - text
        m = re.match(r"^[-*]\s+(.*)$", stripped)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            add_inline_runs(p, m.group(1))
            i += 1
            continue

        # Plain paragraph.
        p = doc.add_paragraph()
        add_inline_runs(p, stripped)
        i += 1

    doc.save(out_path)
    print(f"Wrote {out_path}")


def main():
    if len(sys.argv) < 2:
        print("Usage: python3 scripts/md_to_docx.py <file.md> [file.md ...]")
        sys.exit(1)
    for arg in sys.argv[1:]:
        md_path = Path(arg)
        out_path = md_path.with_suffix(".docx")
        convert(md_path, out_path)


if __name__ == "__main__":
    main()
